import json, re, sys
from pptx import Presentation

POS_TOKEN = r"(?:phr|n|v|vt|vi|adj|adv|num|prep|pron|conj|art|int|aux)\.?"

def extract_lines(path):
    prs = Presentation(path)
    lines = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text
                if t.strip():
                    lines.extend([l.strip() for l in t.split("\n") if l.strip()])
    return [l for l in lines if "大良课堂" not in l and "让优秀成为" not in l]

def truncate_meaning(meaning):
    # 中文释义如果有多个逗号分隔的解释，只保留前两个
    parts = re.split(r"([,，])", meaning)
    result, seg_count = [], 0
    for i, p in enumerate(parts):
        if i % 2 == 0:
            seg_count += 1
            if seg_count > 2:
                break
            result.append(p)
        else:
            if seg_count >= 2:
                break
            result.append(p)
    return "".join(result)

def clean_word(word_part):
    word_part = word_part.strip()
    # strip a trailing POS abbreviation token, e.g. "be excited about phr" -> "be excited about"
    m = re.match(r"^(.*?)\s+" + POS_TOKEN + r"$", word_part)
    if m and m.group(1).strip():
        word_part = m.group(1).strip()
    return word_part.rstrip(".").strip()

def parse_entries(lines):
    entries = []
    failed = []
    prev_num = None
    section = 0
    for l in lines:
        m = re.match(r"^(\d+)\.\s*(.+)$", l)
        if not m:
            continue
        num = int(m.group(1))
        rest = m.group(2).replace("．", ".")
        if prev_num is not None and prev_num > 200 and num <= 3:
            section += 1
        prev_num = num

        # drop a lone 英/美 marker right before a bracket
        rest2 = re.sub(r"\s*(英|美)\s*(?=\[)", " ", rest)

        m2 = re.match(r"^([A-Za-z][A-Za-z\s\.\-']*?)\s*\[[^\]]*\]\s*(.+)$", rest2)
        if m2:
            word = clean_word(m2.group(1))
            meaning = re.sub(r"^\s*" + POS_TOKEN + r"\s*", "", m2.group(2)).strip()
        else:
            cm = re.search(r"[\u4e00-\u9fff]", rest)
            if cm:
                word = clean_word(rest[:cm.start()])
                meaning = rest[cm.start():].strip()
            else:
                word = ""
                meaning = ""
        if not word or not meaning:
            failed.append(l)
            continue
        entries.append({"section": section, "num": num, "word": word, "meaning": meaning})
    return entries, failed

def build(path, prefix, out_path):
    lines = extract_lines(path)
    entries, failed = parse_entries(lines)
    print(f"{path}: parsed {len(entries)} entries, failed {len(failed)}")
    for f in failed[:15]:
        print("  FAIL:", f)

    # 按用户要求：不再是"每36个词一组"，改成固定分 3 个 unit，词平均分到 3 个 unit 里
    UNIT_COUNT = 3
    per_unit = -(-len(entries) // UNIT_COUNT)  # 向上取整
    units = []
    for i in range(0, len(entries), per_unit):
        chunk = entries[i:i + per_unit]
        uid = i // per_unit
        words = []
        for e in chunk:
            wid = f"{prefix}{e['section']}_{e['num']:04d}"
            words.append({"id": wid, "word": e["word"], "meaning": truncate_meaning(e["meaning"])})
        units.append({"unit_index": uid, "words": words})

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({"total_words": len(entries), "units": units}, f, ensure_ascii=False, indent=1)
    print(f"  -> wrote {out_path}: {len(units)} units, sizes {[len(u['words']) for u in units]}")
    return entries

build("C:/Users/wizar/Downloads/初级词汇1-1471&新增.pptx", "j", "junior_words.json")
build("C:/Users/wizar/Downloads/高级词汇1～1404.pptx", "s", "senior_words.json")
